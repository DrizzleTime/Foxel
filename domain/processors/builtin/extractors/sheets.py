import csv
from io import BytesIO, StringIO
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook

from .table import clean_cell, markdown_table
from .text import decode_bytes
from .types import ExtractError, ExtractResult, Section


MAX_SHEET_ROWS = 5000
MAX_SHEET_COLS = 40
_TARGET_TOKENS = 1500


def _estimate_tokens(text: str) -> int:
    """与向量分块相同的估算：非 ASCII 约 1 token，ASCII 约 4 字符 1 token。"""
    tokens = 0
    ascii_run = 0
    for char in text:
        if ord(char) < 128:
            ascii_run += 1
            continue
        if ascii_run:
            tokens += (ascii_run + 3) // 4
            ascii_run = 0
        tokens += 1
    if ascii_run:
        tokens += (ascii_run + 3) // 4
    return tokens


def _limit_rows(rows: Iterable[Any], sheet_name: str) -> tuple[list[list[str]], list[str]]:
    limited: list[list[str]] = []
    warnings: list[str] = []
    truncated_rows = False
    truncated_cols = False
    for index, row in enumerate(rows):
        if index >= MAX_SHEET_ROWS:
            truncated_rows = True
            break
        values = list(row)
        if len(values) > MAX_SHEET_COLS:
            truncated_cols = True
            values = values[:MAX_SHEET_COLS]
        limited.append([clean_cell(value) for value in values])
    if truncated_rows:
        warnings.append(f"工作表 {sheet_name} 仅索引前 {MAX_SHEET_ROWS} 行")
    if truncated_cols:
        warnings.append(f"工作表 {sheet_name} 仅索引前 {MAX_SHEET_COLS} 列")
    return limited, warnings


def sections_from_rows(sheet_name: str, rows: Iterable[Any]) -> tuple[list[Section], list[str]]:
    """每个行窗口重复表头，单独成为一节，避免整表被当成一个原子块。"""
    limited, warnings = _limit_rows(rows, sheet_name)
    cleaned = [row for row in limited if any(cell for cell in row)]
    if not cleaned:
        return [], warnings

    header = cleaned[0]
    body = cleaned[1:]
    header_tokens = _estimate_tokens(markdown_table([header]))
    windows: list[list[list[str]]] = []
    current: list[list[str]] = []
    current_tokens = header_tokens

    for row in body:
        row_tokens = _estimate_tokens("| " + " | ".join(row) + " |")
        if current and current_tokens + row_tokens > _TARGET_TOKENS:
            windows.append(current)
            current = [row]
            current_tokens = header_tokens + row_tokens
        else:
            current.append(row)
            current_tokens += row_tokens
    if current or not body:
        windows.append(current)

    sections: list[Section] = []
    for window in windows:
        text = markdown_table([header, *window])
        if not text:
            continue
        sections.append(Section(
            heading=sheet_name,
            text=text,
            kind="sheet",
            locator={"sheet": sheet_name},
        ))
    return sections, warnings


def _sections_from_tables(tables: list[tuple[str, Iterable[Any]]]) -> ExtractResult:
    sections: list[Section] = []
    warnings: list[str] = []
    for name, rows in tables:
        sheet_sections, sheet_warnings = sections_from_rows(name, rows)
        sections.extend(sheet_sections)
        warnings.extend(sheet_warnings)
    if not sections:
        raise ExtractError("表格没有可索引的内容")
    return ExtractResult(sections=sections, warnings=warnings)


def extract_xlsx(data: bytes) -> ExtractResult:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractError("表格无法解析") from exc
    sections: list[Section] = []
    warnings: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            sheet_sections, sheet_warnings = sections_from_rows(
                worksheet.title or "Sheet",
                worksheet.iter_rows(values_only=True),
            )
            sections.extend(sheet_sections)
            warnings.extend(sheet_warnings)
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("表格无法解析") from exc
    finally:
        workbook.close()
    if not sections:
        raise ExtractError("表格没有可索引的内容")
    return ExtractResult(sections=sections, warnings=warnings)


def _xls_rows(book, sheet) -> list[list[Any]]:
    rows: list[list[Any]] = []
    row_count = min(sheet.nrows, MAX_SHEET_ROWS + 1)
    col_count = min(sheet.ncols, MAX_SHEET_COLS + 1)
    for row_index in range(row_count):
        row: list[Any] = []
        for col_index in range(col_count):
            cell_type = sheet.cell_type(row_index, col_index)
            value = sheet.cell_value(row_index, col_index)
            if cell_type == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                except Exception:
                    pass
            elif cell_type == xlrd.XL_CELL_ERROR:
                value = ""
            elif cell_type == xlrd.XL_CELL_BOOLEAN:
                value = bool(value)
            row.append(value)
        rows.append(row)
    return rows


def extract_xls(data: bytes) -> ExtractResult:
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ExtractError("表格无法解析") from exc
    tables = [(sheet.name or "Sheet", _xls_rows(book, sheet)) for sheet in book.sheets()]
    return _sections_from_tables(tables)


def extract_csv(data: bytes, sheet_name: str) -> ExtractResult:
    try:
        text = decode_bytes(data)
    except ExtractError:
        raise
    try:
        rows = list(csv.reader(StringIO(text)))
    except csv.Error as exc:
        raise ExtractError("表格无法解析") from exc
    return _sections_from_tables([(sheet_name or "CSV", rows)])
