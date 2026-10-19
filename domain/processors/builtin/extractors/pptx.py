from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .table import clean_cell, markdown_table
from .types import ExtractError, ExtractResult, Section


def _table_text(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [clean_cell(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return markdown_table(rows)


def _is_title_shape(shape, title_shape) -> bool:
    if title_shape is None:
        return False
    if shape is title_shape:
        return True
    return getattr(shape, "_element", None) is getattr(title_shape, "_element", None)


def _walk_shapes(shapes, title_shape, parts: list[str]) -> None:
    for shape in shapes:
        if _is_title_shape(shape, title_shape):
            continue
        shape_type = getattr(shape, "shape_type", None)
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            _walk_shapes(shape.shapes, title_shape, parts)
            continue
        if getattr(shape, "has_table", False):
            table_md = _table_text(shape.table)
            if table_md:
                parts.append(table_md)
            continue
        if getattr(shape, "has_text_frame", False):
            text = "\n".join(line.strip() for line in shape.text_frame.text.splitlines() if line.strip())
            if text:
                parts.append(text)


def _notes_text(slide) -> str:
    if not getattr(slide, "has_notes_slide", False):
        return ""
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return ""
    return "\n".join(line.strip() for line in frame.text.splitlines() if line.strip())


def extract_pptx(data: bytes) -> ExtractResult:
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise ExtractError("演示文稿无法解析") from exc

    sections: list[Section] = []
    for index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            title = " ".join(title_shape.text.split())
        parts: list[str] = []
        _walk_shapes(slide.shapes, title_shape, parts)
        notes = _notes_text(slide)
        body = "\n\n".join(parts).strip()
        if notes and notes not in body:
            body = f"{body}\n\n{notes}".strip() if body else notes
        if not body and title:
            body = title
        if not body:
            continue
        heading = f"Slide {index}: {title}" if title else f"Slide {index}"
        sections.append(Section(
            heading=heading,
            text=body,
            kind="slide",
            locator={"slide": index},
        ))

    if not sections:
        raise ExtractError("演示文稿没有可索引的文本")
    return ExtractResult(sections=sections, warnings=[])
